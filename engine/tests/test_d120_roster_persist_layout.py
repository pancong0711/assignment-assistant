"""D120：workspace 根语义收口 + roster 自动落盘 + 目录错位迁移 + 作业纸落盘。

触发：roster 被写进 `<ws>/kb/classes/`（引擎只读 `<ws>/<class_dir>/roster/`）
→ AI 批阅报「名单 0 人」。
"""
import json
from pathlib import Path

import pytest
from click.testing import CliRunner
from openpyxl import load_workbook

from assist import workspace_layout as WL
from assist.cli import cli
from assist.files import roster_io, taskpads
from assist.grading import inputs as _inputs


# ---------------- roster 读写（路径由引擎决定） ----------------

def test_write_roster_goes_to_engine_path(tmp_path):
    """写入位置必须与 `grading/inputs.roster_file()` 读取位置**完全一致**。"""
    ws = tmp_path / "ws"
    rep = roster_io.write_roster(ws, "classes/化工25", [
        {"name": "张三", "number": "01", "class": "化工25", "tag": "distinguish"},
        {"name": "李四", "number": "02", "class": "化工25", "tag": "copy"},
    ])
    assert rep["ok"] and rep["count"] == 2
    # 引擎读得到（这是关键：写/读同口径）
    fn = _inputs.roster_file(ws, "classes/化工25")
    assert fn is not None and fn.is_file()
    rows = _inputs.students_with_tags(ws, "classes/化工25")
    assert {r["name"] for r in rows} == {"张三", "李四"}
    assert {r["tag"] for r in rows} == {"distinguish", "copy"}


def test_write_roster_rejects_bad_class_dir(tmp_path):
    for bad in ("", "..", "/abs/path", "classes/../../etc"):
        rep = roster_io.write_roster(tmp_path, bad, [{"name": "张三"}])
        assert rep["ok"] is False, f"应拒绝 {bad!r}"


def test_write_roster_backs_up_existing(tmp_path):
    ws = tmp_path / "ws"
    roster_io.write_roster(ws, "classes/A", [{"name": "张三"}])
    rep = roster_io.write_roster(ws, "classes/A", [{"name": "李四"}])
    assert rep["ok"]
    assert (ws / "classes/A/roster/roster.xlsx").is_file()
    olds = list((ws / "classes/A/roster").glob("roster.xlsx.old-*"))
    assert olds, "覆盖前应留备份"


def test_empty_roster_not_written(tmp_path):
    rep = roster_io.write_roster(tmp_path, "classes/A", [])
    assert rep["ok"] is False and "为空" in rep["error"]


def test_read_roster_roundtrip(tmp_path):
    ws = tmp_path / "ws"
    roster_io.write_roster(ws, "classes/A", [{"name": "张三", "number": "01", "class": "A", "tag": "t"}])
    got = roster_io.read_roster_file(ws, "classes/A")
    assert got["ok"] and got["count"] == 1 and got["students"][0]["tag"] == "t"


# ---------------- 目录错位迁移（kb/classes → classes） ----------------

def test_inspect_detects_kb_classes(tmp_path):
    ws = tmp_path / "ws"
    (ws / "kb" / "classes" / "化工25" / "roster").mkdir(parents=True)
    (ws / "kb" / "classes" / "化工25" / "roster" / "roster.xlsx").write_bytes(b"x")
    inv = WL.inspect(ws)
    assert inv["summary"]["violations"] == 1
    assert inv["summary"]["files"] == 1
    assert inv["migratable"][0]["to"].endswith("classes/化工25/roster/roster.xlsx")


def test_fix_layout_dry_run_then_apply(tmp_path):
    ws = tmp_path / "ws"
    src = ws / "kb" / "classes" / "化工25" / "roster" / "roster.xlsx"
    src.parent.mkdir(parents=True)
    src.write_bytes(b"data")
    dry = WL.fix_layout(ws, apply=False)
    assert dry["summary"]["moved"] == 1
    assert not (ws / "classes" / "化工25" / "roster" / "roster.xlsx").exists(), "演练不应落盘"
    rep = WL.fix_layout(ws, apply=True)
    dst = ws / "classes" / "化工25" / "roster" / "roster.xlsx"
    assert dst.is_file() and dst.read_bytes() == b"data"
    assert src.is_file(), "**不删源**"
    # 幂等
    rep2 = WL.fix_layout(ws, apply=True)
    assert rep2["summary"]["moved"] == 0 and rep2["summary"]["skipped"] == 1


def test_fix_layout_does_not_overwrite(tmp_path):
    ws = tmp_path / "ws"
    src = ws / "kb" / "classes" / "A" / "roster" / "roster.xlsx"
    src.parent.mkdir(parents=True)
    src.write_bytes(b"old")
    dst = ws / "classes" / "A" / "roster" / "roster.xlsx"
    dst.parent.mkdir(parents=True)
    dst.write_bytes(b"new")
    rep = WL.fix_layout(ws, apply=True)
    assert rep["summary"]["moved"] == 0
    assert dst.read_bytes() == b"new", "不得覆盖已有文件"


# ---------------- 作业纸落盘 ----------------

def test_taskpad_save_and_list(tmp_path):
    ws = tmp_path / "ws"
    rep = taskpads.save_taskpad(ws, {"id": "p1", "items": [], "layout": {}})
    assert rep["ok"]
    p = ws / "tasks" / "p1.taskpad.json"
    assert p.is_file()
    assert json.loads(p.read_text(encoding="utf-8"))["id"] == "p1"
    assert [x["id"] for x in taskpads.list_taskpads(ws)] == ["p1"]


def test_taskpad_id_sanitized(tmp_path):
    ws = tmp_path / "ws"
    rep = taskpads.save_taskpad(ws, {"id": "../evil/../x", "items": []})
    assert rep["ok"]
    assert ".." not in Path(rep["path"]).name


# ---------------- CLI 等价（教师要求 CLI 能重复 PWA 功能） ----------------

def test_cli_workspace_layout_and_fix(tmp_path):
    ws = tmp_path / "ws"
    src = ws / "kb" / "classes" / "A" / "roster" / "roster.xlsx"
    src.parent.mkdir(parents=True)
    src.write_bytes(b"x")
    r = CliRunner().invoke(cli, ["workspace", "layout", "-w", str(ws)])
    assert r.exit_code == 0 and '"violations"' in r.output
    r2 = CliRunner().invoke(cli, ["workspace", "fix-layout", "-w", str(ws)])
    assert r2.exit_code == 0 and "演练" in r2.output
    r3 = CliRunner().invoke(cli, ["workspace", "fix-layout", "--apply", "-w", str(ws)])
    assert r3.exit_code == 0
    assert (ws / "classes" / "A" / "roster" / "roster.xlsx").is_file()


def test_cli_roster_save_and_show(tmp_path):
    ws = tmp_path / "ws"
    js = tmp_path / "roster.json"
    js.write_text(json.dumps([{"name": "张三", "number": "01", "class": "A", "tag": "t"}],
                             ensure_ascii=False), encoding="utf-8")
    r = CliRunner().invoke(cli, ["workspace", "roster-save", "--class-dir", "classes/A",
                                 "--from-json", str(js), "-w", str(ws)])
    assert r.exit_code == 0, r.output
    assert (ws / "classes" / "A" / "roster" / "roster.xlsx").is_file()
    r2 = CliRunner().invoke(cli, ["workspace", "roster-show", "--class-dir", "classes/A",
                                  "-w", str(ws)])
    assert r2.exit_code == 0 and "张三" in r2.output


# ---------------- doctor 检查项 ----------------

def test_doctor_detects_layout_and_offers_api_fix(tmp_path):
    import assist.serve as S
    ws = tmp_path / "ws"
    (ws / "kb" / "classes" / "A" / "roster").mkdir(parents=True)
    (ws / "kb" / "classes" / "A" / "roster" / "roster.xlsx").write_bytes(b"x")
    checks = {c["id"]: c for c in S._doctor_checks(ws)}
    c = checks["workspace_layout"]
    assert c["status"] == "yellow"
    assert c["fix"].get("type") == "api"
    assert c["fix"].get("endpoint") == "/workspace/fix-layout"


# ---------------- 审核补充（D120 复审发现） ----------------

def test_roster_concurrent_writes_keep_backups(tmp_path):
    """并发写同一份 roster：不丢内容、备份不互相覆盖（加锁前 8 并发只剩 1 个备份）。"""
    import threading

    from openpyxl import load_workbook
    ws = tmp_path / "ws"
    errs: list[str] = []

    def w(i: int) -> None:
        try:
            roster_io.write_roster(ws, "classes/A", [{"name": f"学生{i}"}])
        except Exception as exc:  # noqa: BLE001
            errs.append(str(exc))

    ts = [threading.Thread(target=w, args=(i,)) for i in range(8)]
    [t.start() for t in ts]
    [t.join() for t in ts]
    assert not errs
    p = ws / "classes/A/roster/roster.xlsx"
    wb = load_workbook(p, read_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    wb.close()
    assert rows[0] == ("name", "number", "class", "tag")
    assert len(rows) - 1 == 1, "最终应是一份完整名单（不是拼接/半截）"
    assert len(list((ws / "classes/A/roster").glob("roster.xlsx.old-*"))) >= 6, \
        "并发的历史备份不应互相覆盖"
    assert not list((ws / "classes/A/roster").glob("*.tmp")), "不应残留临时文件"


def test_fix_layout_ignores_normal_kb_content(tmp_path):
    """正常 kb 内容（fig/、solutions/、xlsx）不得被误判/误搬。"""
    ws = tmp_path / "ws"
    (ws / "kb" / "fig").mkdir(parents=True)
    (ws / "kb" / "problems.xlsx").write_bytes(b"x")
    (ws / "kb" / "fig" / "a.png").write_bytes(b"p")
    (ws / "kb" / "solutions").mkdir()
    (ws / "kb" / "solutions" / "c1.json").write_bytes(b"{}")
    inv = WL.inspect(ws)
    assert inv["summary"] == {"violations": 0, "files": 0}
    assert WL.fix_layout(ws, apply=True)["summary"]["moved"] == 0
    assert (ws / "kb" / "fig" / "a.png").is_file()


def test_taskpad_too_large_rejected(tmp_path):
    """作业纸过大应拒绝写盘（防把磁盘写爆），但正常的不受影响。"""
    ws = tmp_path / "ws"
    assert taskpads.save_taskpad(ws, {"id": "ok", "items": []})["ok"]
    big = {"id": "big", "items": [{"x": "y" * 1000}] * 9000}
    rep = taskpads.save_taskpad(ws, big)
    assert rep["ok"] is False and "过大" in rep["error"]


def test_roster_write_read_path_identical_to_engine(tmp_path):
    """写盘路径与 `grading/inputs.roster_file()` 读取路径必须**完全一致**（防再次漂移）。"""
    ws = tmp_path / "ws"
    roster_io.write_roster(ws, "classes/化工25", [{"name": "张三", "tag": "t"}])
    written = ws / "classes/化工25/roster/roster.xlsx"
    read_by = _inputs.roster_file(ws, "classes/化工25")
    assert written.resolve() == read_by.resolve()


def test_inspect_always_has_summary(tmp_path):
    """`inspect()` 无论 kb/ 是否存在都必须带 summary（doctor/调用方假设该字段）。"""
    # kb/ 不存在
    assert WL.inspect(tmp_path)["summary"] == {"violations": 0, "files": 0}
    # kb/ 存在但无错位
    (tmp_path / "kb").mkdir()
    assert WL.inspect(tmp_path)["summary"] == {"violations": 0, "files": 0}
