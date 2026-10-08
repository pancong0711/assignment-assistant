"""D63 T2 回归：学习通统计导出解析 + 插班标注 + 交叉核验（合成 xlsx，纯本地）。"""

from pathlib import Path

from openpyxl import Workbook

from assist.xxt.master_roster import (annotate_years, build, cross_check,
                                      parse_export_banner, parse_stat_export,
                                      score_sheet_names)


def _mk_export(path: Path, students: list[tuple[str, str, str]]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(["课程：大学物理C1", None, None])
    ws.append(["学生姓名", "学号/工号", "函授站", "层次", "专业", "作业(35%)"])
    for n, num, cls in students:
        ws.append([n, num, "某学院", "某专业(本)", cls, "30"])
    ws.append(["ideal", "", "", "", "", ""])
    wb.save(path)
    return path


def _mk_scores(path: Path, names: list[str]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "化工24-分析"
    ws.append(["", "1", "2"])
    ws.append(["姓名", "得分", "失分"])
    for n in names:
        ws.append([n, "0", "-2"])
    ws.append(["ideal", "", ""])
    ws.append(["rate", "", ""])
    wb.save(path)
    return path


def test_parse_export_and_years(tmp_path):
    fn = _mk_export(tmp_path / "e.xlsx",
                    [("张三", "2024438815101", "化工241"),
                     ("李四", "2024438815102", "化工241"),
                     ("王五", "2023438815999", "化工241")])
    out = parse_stat_export(fn)
    assert out["banner"] == {}  # 合成数据无班级 banner 行
    assert len(out["students"]) == 3
    assert out["students"][0]["class"] == "化工241"
    assert out["other_rows"] == ["ideal"]
    years = annotate_years(out["students"])
    assert dict(years) == {"2024": 2, "2023": 1}
    assert out["students"][2]["note"] == "插班(2023级)"


def test_parse_banner(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["课程：大学物理C1    班级：潘聪-化工24    任课教师：戈迪  导出时间：2025-06-21 16:50:46"])
    wb.save(path := tmp_path / "b.xlsx")
    assert parse_export_banner(list(ws.iter_rows(values_only=True)))["class"] == "潘聪-化工24"


def test_scores_and_cross_check(tmp_path):
    sp = _mk_scores(tmp_path / "s.xlsx", ["张三", "李四", "赵六"])
    assert score_sheet_names(sp, "化工24-分析") == {"张三", "李四", "赵六"}
    master = [{"name": "张三"}, {"name": "李四"}, {"name": "王五"}]
    cc = cross_check(master, {"张三", "李四", "赵六"})
    assert cc["both"] == 2 and cc["only_master"] == ["王五"] and cc["only_scores"] == ["赵六"]


def test_build(tmp_path):
    e1 = _mk_export(tmp_path / "a.xlsx", [("张三", "2024438815101", "化工241"),
                                          ("王五", "2023438815999", "化工241")])
    e2 = _mk_export(tmp_path / "b.xlsx", [("李四", "2024438905101", "环境241")])
    s1 = _mk_scores(tmp_path / "s.xlsx", ["张三"])
    out = tmp_path / ".out" / "master.xlsx"
    summary = build({"化工24": e1, "环境24": e2},
                    cross_sources={"化工24": (s1, "化工24-分析")}, out_xlsx=out)
    assert summary["total"] == 3
    assert summary["classes"]["化工24"]["cross"]["only_scores"] == []
    assert out.exists()
    from openpyxl import load_workbook
    wb = load_workbook(out, read_only=True)
    rows = list(wb["namelist_df"].iter_rows(values_only=True))
    wb.close()
    assert rows[0][:3] == ("name", "number", "class")
    assert len(rows) == 4  # 表头 + 3 行
