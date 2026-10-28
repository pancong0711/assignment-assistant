"""D110：答案版本库按 (kb, chap, qid) 作用域隔离 + 防"误刷全库"。

背景 bug：
1) 不同章节存在同号题（chap04/c1、chap05/c1），答案文件原先只按 qid 命名 →
   两章写进同一个 c1.json，读回时两章显示同一份答案。
2) 前端"题号"多选框留空 = 整章/整库；若章节也没选就变成**整个题库**作答，
   用户没点全库却刷了全库。
"""
import json
import tempfile
import urllib.error
import urllib.request
import threading
from http.server import ThreadingHTTPServer
from pathlib import Path

import pytest
from click.testing import CliRunner
from openpyxl import Workbook

from assist.cli import cli
from assist.files import solutions as sol
from assist.grading import solve


def _mk_kb(ws: Path) -> None:
    kb = ws / "kb"
    kb.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    for title, ids in (("chap04", ["c1", "c2"]), ("chap05", ["c1", "c3"])):
        sh = wb.create_sheet(title)
        sh.append(["id", "content", "img_path", "page", "related", "type", "solution", "note"])
        for i in ids:
            sh.append([i, f"{title} 的题 {i}", "", 1, "", "计算", "", ""])
    wb.save(kb / "problems.xlsx")


# ---------------- 跨章同号隔离 ----------------

def test_same_qid_in_different_chapters_are_isolated(tmp_path):
    ws = tmp_path / "ws"; _mk_kb(ws)
    sol.add_version(ws, "chap04-c1", "第四章答案", kb="problems", chap="chap04", pin=True)
    sol.add_version(ws, "chap05-c1", "第五章答案", kb="problems", chap="chap05", pin=True)

    assert sol.effective_solution(ws, "chap04-c1", kb="problems", chap="chap04")["text"] == "第四章答案"
    assert sol.effective_solution(ws, "chap05-c1", kb="problems", chap="chap05")["text"] == "第五章答案"
    # 两个独立文件
    files = sorted(p.name for p in (ws / "kb" / "solutions").glob("*.json"))
    assert files == ["problems__chap04__chap04-c1.json",
                     "problems__chap05__chap05-c1.json"]


def test_scope_id_format():
    assert sol.scope_id("problems", "chap04", "c1") == "problems__chap04__c1"
    assert sol.scope_id("problems", "", "c1") == "problems__c1"
    assert sol.scope_id("", "", "c1") == "c1"


def test_effective_by_key_and_ambiguous_flag(tmp_path):
    ws = tmp_path / "ws"; _mk_kb(ws)
    sol.add_version(ws, "c1", "第四章", kb="problems", chap="chap04", pin=True)
    sol.add_version(ws, "c1", "第五章", kb="problems", chap="chap05", pin=True)
    by_key = sol.effective_by_key(ws)
    assert by_key["problems__chap04__c1"]["text"] == "第四章"
    assert by_key["problems__chap05__c1"]["text"] == "第五章"
    # 按 qid 的旧索引必须标记 ambiguous（避免前端盲用）
    m = sol.effective_map(ws)
    assert m["c1"].get("ambiguous") is True


def test_legacy_qid_file_is_migrated_and_not_leaked(tmp_path):
    """旧的无章节文件 c1.json 在首次按章节写入后应被标记迁移，不再串到别的章节。"""
    ws = tmp_path / "ws"; _mk_kb(ws)
    sol.add_version(ws, "c1", "旧人工答案", source="human", pin=True)   # 旧式：无作用域
    sol.add_version(ws, "c1", "第四章AI", kb="problems", chap="chap04", pin=True)

    assert sol.effective_solution(ws, "c1", kb="problems", chap="chap04")["text"] == "第四章AI"
    # 其它章节不应读到第四章的答案
    assert sol.effective_solution(ws, "c1", kb="problems", chap="chap05")["text"] == ""
    legacy = json.loads((ws / "kb" / "solutions" / "c1.json").read_text(encoding="utf-8"))
    assert legacy.get("migratedTo", "").endswith("problems__chap04__c1.json")


def test_pin_and_remove_respect_scope(tmp_path):
    ws = tmp_path / "ws"; _mk_kb(ws)
    sol.add_version(ws, "c1", "四-1", kb="problems", chap="chap04")
    sol.add_version(ws, "c1", "四-2", kb="problems", chap="chap04")
    sol.add_version(ws, "c1", "五-1", kb="problems", chap="chap05")
    # 置顶第四章的 v2 不应影响第五章
    assert sol.pin_version(ws, "c1", "v2", kb="problems", chap="chap04")["ok"]
    assert sol.effective_solution(ws, "c1", kb="problems", chap="chap04")["text"] == "四-2"
    assert sol.effective_solution(ws, "c1", kb="problems", chap="chap05")["text"] == "五-1"
    # 删除第五章的 v1 不影响第四章
    assert sol.remove_version(ws, "c1", "v1", kb="problems", chap="chap05")["ok"]
    assert sol.effective_solution(ws, "c1", kb="problems", chap="chap04")["text"] == "四-2"


# ---------------- 防"误刷全库" ----------------

def test_list_targets_scope():
    ws = Path(tempfile.mkdtemp()) / "ws"; _mk_kb(ws)
    only_chap = solve.list_targets(ws, "problems", "chap04", [])
    # D111：题库 id 全局唯一，读入即带章节前缀
    assert {t["id"] for t in only_chap} == {"chap04-c1", "chap04-c2"}
    # 传裸题号仍可命中（历史作业纸兼容）
    picked = solve.list_targets(ws, "problems", "chap04", ["c1"])
    assert [(t["chap"], t["id"]) for t in picked] == [("chap04", "chap04-c1")]
    whole = solve.list_targets(ws, "problems", "", [])
    assert len(whole) == 4, "无范围=整库（调用方必须先确认）"


def test_cli_ai_solve_refuses_implicit_whole_library(tmp_path):
    ws = tmp_path / "ws"; _mk_kb(ws)
    r = CliRunner().invoke(cli, ["kb", "ai-solve", "--kb-kind", "problems", "-w", str(ws)])
    assert r.exit_code != 0
    assert "未指定范围" in r.output
    r2 = CliRunner().invoke(cli, ["kb", "ai-solve", "--kb-kind", "problems",
                                  "--chap", "chap04", "-w", str(ws)])
    assert r2.exit_code == 0, r2.output
    assert json.loads(r2.output.split("\n（dry-run")[0])["dryRun"] is True


def test_cli_ai_solve_all_flag_allows_whole_library(tmp_path):
    ws = tmp_path / "ws"; _mk_kb(ws)
    r = CliRunner().invoke(cli, ["kb", "ai-solve", "--kb-kind", "problems", "--all",
                                 "-w", str(ws)])
    assert r.exit_code == 0, r.output
    assert json.loads(r.output.split("\n（dry-run")[0])["planned"] == 4


def test_serve_ai_solve_requires_scope(tmp_path):
    import assist.serve as S

    ws = tmp_path / "ws"; _mk_kb(ws)
    S._set_serve_ws(ws); S.Handler.ws = ws.resolve(); S.Handler.token = ""
    httpd = ThreadingHTTPServer(("127.0.0.1", 0), S.Handler)
    port = httpd.server_address[1]
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    try:
        base = f"http://127.0.0.1:{port}"
        body = json.dumps({"kb": "problems", "ids": []}).encode()
        req = urllib.request.Request(base + "/kb/ai-solve", data=body,
                                     headers={"Content-Type": "application/json"}, method="POST")
        with pytest.raises(urllib.error.HTTPError) as ei:
            urllib.request.urlopen(req)
        assert ei.value.code == 400
        assert "未指定范围" in json.loads(ei.value.read())["error"]
    finally:
        httpd.shutdown(); S._SERVE_WS = None
