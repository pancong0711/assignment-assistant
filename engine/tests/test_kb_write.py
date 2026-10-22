"""B4：openpyxl 样式保留写回。"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from assist.files.kb_io import write_chapters_preserving


def test_write_chapters_preserving_keeps_cell_style(tmp_path: Path):
    xlsx = tmp_path / "problems.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "chap04"
    ws.append(["id", "content", "img_path", "page", "related", "type", "solution", "note"])
    ws.append(["old-1", "旧题", "", 1, "", "计算", "旧答案", "旧备注"])
    ws["A1"].font = Font(bold=True)
    ws["B2"].fill = PatternFill("solid", fgColor="FFFF00")
    wb.save(xlsx)

    write_chapters_preserving(xlsx, {
        "chap04": [{"id": "new-1", "content": "新题", "page": 2,
                    "type": "计算", "solution": "新答案"}]
    })
    wb2 = Workbook()
    wb2 = __import__("openpyxl").load_workbook(xlsx)
    ws2 = wb2["chap04"]
    assert ws2["A1"].value == "id"
    assert ws2["A1"].font.bold is True          # 表头样式保留
    assert ws2["B2"].value == "新题"
    assert ws2["B2"].fill.fgColor.rgb == "00FFFF00"  # 单元格填充保留
    assert ws2.max_row == 2                     # 旧第 2 行后的数据被清掉
    wb2.close()


def test_validate_sheet_name():
    from assist.files.kb_io import validate_sheet_name
    assert validate_sheet_name("chap04") == ""
    assert validate_sheet_name("章-01_复习") == ""
    assert validate_sheet_name("") != ""
    assert validate_sheet_name("a" * 32) != ""
    assert validate_sheet_name("ch[1]") != ""
    assert validate_sheet_name("ch/1") != ""


def test_write_chapters_preserving_rejects_bad_sheet(tmp_path: Path):
    import pytest
    from assist.files.kb_io import write_chapters_preserving
    with pytest.raises(ValueError):
        write_chapters_preserving(tmp_path / "problems.xlsx", {"ch[1]": []})


def test_cli_kb_write_rejects_bad_sheet(tmp_path: Path):
    import json
    from click.testing import CliRunner
    from assist.cli import cli
    ws = tmp_path / "ws"
    (ws / "kb").mkdir(parents=True)
    payload = tmp_path / "in.json"
    payload.write_text(json.dumps({"ch[1]": [{"id": "x", "content": "题"}]}), encoding="utf-8")
    r = CliRunner().invoke(cli, ["kb", "write", "--kind", "problems",
                                 "--input", str(payload), "-w", str(ws)])
    assert r.exit_code != 0
    assert "章节名非法" in r.output


def test_cli_kb_stats_empty_hint_mentions_real_entry(tmp_path):
    """P0：空题库提示不得再指向不存在的 `assist kb init`。"""
    from click.testing import CliRunner
    from assist.cli import cli
    ws = tmp_path / "ws"
    (ws / "kb").mkdir(parents=True)
    r = CliRunner().invoke(cli, ["kb", "stats", "-w", str(ws)])
    assert r.exit_code == 0, r.output
    assert "kb init" not in r.output
    assert "kb_init" in r.output
