"""D111：题库题号全局唯一（自动章节前缀 `<chap>-<id>`）。

背景：题号只在章内唯一（真实题库里 `01` 在 9 个章节重复），导致答案版本库串味、
批阅报告错配。约定读入即补 `<chap>-` 前缀（幂等），并兼容历史裸 id 的作业纸。
"""
import json
from pathlib import Path

from click.testing import CliRunner
from openpyxl import Workbook, load_workbook

from assist.cli import cli
from assist.files import solutions as sol
from assist.files.kb_io import (find_duplicate_ids, normalize_question_ids,
                                read_kb, scope_question_id)
from assist.grading import solve
from assist.paper.task import item_id_matches, resolve_items


def _mk_kb(ws: Path, chaps=("chap04", "chap07")) -> None:
    kb = ws / "kb"; kb.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    for name in chaps:
        sh = wb.create_sheet(name)
        sh.append(["id", "content", "img_path", "page", "related", "type", "solution", "note"])
        sh.append(["01", f"{name} 第一题", "", 1, "", "计算", "", ""])
        sh.append(["02", f"{name} 第二题", "", 1, "", "计算", "", ""])
    wb.save(kb / "problems.xlsx")


# ---------------- 前缀函数 ----------------

def test_scope_question_id_basic_and_idempotent():
    assert scope_question_id("chap04", "01") == "chap04-01"
    assert scope_question_id("chap04", "chap04-01") == "chap04-01", "已带前缀不得重复加"
    assert scope_question_id("chap04", "demo-04-1") == "chap04-demo-04-1"
    assert scope_question_id("", "01") == "01"
    assert scope_question_id("chap04", "") == ""


def test_scope_question_id_sanitizes():
    assert scope_question_id("ch/04", "x") == "ch-04-x"
    assert scope_question_id("chap04", "a/b") == "chap04-a-b"


def test_normalize_is_idempotent():
    ch = {"chap04": [{"id": "01"}, {"id": "02"}]}
    once, changed = normalize_question_ids(ch)
    assert changed == [{"chap": "chap04", "old": "01", "new": "chap04-01"},
                       {"chap": "chap04", "old": "02", "new": "chap04-02"}]
    twice, changed2 = normalize_question_ids(once)
    assert twice == once and changed2 == []


def test_duplicates_detected_same_sheet_and_cross_sheet():
    # 同 sheet 内重复
    assert list(find_duplicate_ids({"a": [{"id": "x"}, {"id": "x"}]})) == ["a-x"]
    # 跨 sheet 同号 → 加前缀后不重复
    assert find_duplicate_ids({"chap04": [{"id": "01"}], "chap07": [{"id": "01"}]}) == {}
    # 已含本章前缀的另一写法 → 撞车
    assert list(find_duplicate_ids({"chap07": [{"id": "01"}, {"id": "chap07-01"}]})) == ["chap07-01"]


# ---------------- 读入即全局唯一 ----------------

def test_read_kb_prefixes_ids(tmp_path):
    ws = tmp_path / "ws"; _mk_kb(ws)
    kb = read_kb(ws / "kb")
    assert [e["id"] for e in kb["problems"]["chap04"]] == ["chap04-01", "chap04-02"]
    assert [e["id"] for e in kb["problems"]["chap07"]] == ["chap07-01", "chap07-02"]


def test_cross_chapter_same_number_gets_distinct_answers(tmp_path):
    ws = tmp_path / "ws"; _mk_kb(ws)
    sol.add_version(ws, "chap04-01", "第四章答案", kb="problems", chap="chap04", pin=True)
    sol.add_version(ws, "chap07-01", "第七章答案", kb="problems", chap="chap07", pin=True)
    assert sol.effective_solution(ws, "chap04-01", kb="problems", chap="chap04")["text"] == "第四章答案"
    assert sol.effective_solution(ws, "chap07-01", kb="problems", chap="chap07")["text"] == "第七章答案"


# ---------------- 历史作业纸兼容（裸题号） ----------------

def test_item_id_matches_accepts_bare_and_global():
    assert item_id_matches("chap04", ["01"], "chap04-01")
    assert item_id_matches("chap04", ["chap04-01"], "chap04-01")
    assert not item_id_matches("chap04", ["02"], "chap04-01")
    assert not item_id_matches("chap04", [""], "chap04-01")


def test_resolve_items_accepts_legacy_bare_ids(tmp_path):
    ws = tmp_path / "ws"; _mk_kb(ws)
    kb = read_kb(ws / "kb")
    legacy = [{"kb": "problems", "chap": "chap04", "ids": ["01"], "tag": "t"}]
    out = resolve_items(kb, legacy, ws / "kb")
    assert len(out) == 1 and "第一题" in out[0][0], "历史裸题号必须仍能命中"
    modern = [{"kb": "problems", "chap": "chap04", "ids": ["chap04-01"], "tag": "t"}]
    assert len(resolve_items(kb, modern, ws / "kb")) == 1


def test_list_targets_accepts_bare_ids(tmp_path):
    ws = tmp_path / "ws"; _mk_kb(ws)
    picked = solve.list_targets(ws, "problems", "chap04", ["01"])
    assert [(t["chap"], t["id"]) for t in picked] == [("chap04", "chap04-01")]


# ---------------- CLI normalize ----------------

def test_cli_normalize_dry_run_then_write(tmp_path):
    ws = tmp_path / "ws"; _mk_kb(ws)
    r = CliRunner().invoke(cli, ["kb", "normalize", "-w", str(ws)])
    assert r.exit_code == 0, r.output
    rep = json.loads(r.output)
    assert rep["write"] is False
    assert rep["kinds"]["problems"]["sample_ids"] == ["chap04-01", "chap04-02"]

    # dry-run 不落盘
    wb = load_workbook(ws / "kb" / "problems.xlsx", read_only=True)
    assert [row[0] for row in list(wb["chap04"].iter_rows(values_only=True))[1:]] == ["01", "02"]
    wb.close()

    r2 = CliRunner().invoke(cli, ["kb", "normalize", "--write", "-w", str(ws)])
    assert r2.exit_code == 0, r2.output
    wb = load_workbook(ws / "kb" / "problems.xlsx", read_only=True)
    assert [row[0] for row in list(wb["chap04"].iter_rows(values_only=True))[1:]] == ["chap04-01", "chap04-02"]
    assert [row[0] for row in list(wb["chap07"].iter_rows(values_only=True))[1:]] == ["chap07-01", "chap07-02"]
    wb.close()
