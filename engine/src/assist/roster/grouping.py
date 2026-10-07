"""M5 成绩管理（docs/05-D17/D18）——算法迁移自
_legacy/2603paperDesign/src/paperdesign/student/student.py：

- _sortByScore/_tag/_specialTag 语义逐条保留（本实现去 pandas，纯标准库+openpyxl）；
- 9 种 tag 见 docs/05-D17；punish 兼容旧名 penalty/penalty2；
- translation 组的"随机点选"改为可复现的等距选点（CLI 与 PWA 前端口径一致）；
- 比例组顺序：distinguish→innovation→summary→qa→copy→copyOnly（与旧 default 一致）。
"""

from pathlib import Path

from loguru import logger
from openpyxl import load_workbook

DEFAULT_GROUP_CFG = [
    {"group_name": "distinguish", "group_ratio": 0.05},
    {"group_name": "innovation", "group_ratio": 0.10},
    {"group_name": "summary", "group_ratio": 0.10},
    {"group_name": "qa", "group_ratio": 0.25},
    {"group_name": "copy", "group_ratio": 0.35},
    {"group_name": "copyOnly", "group_ratio": 0.15},
    {"group_name": "translation", "group_ratio": 0.10},
]

RANDOM_GROUP = "translation"          # 迁移注释：旧版以 np.random 点选
PENALTY_TAGS = {"punish", "penalty", "penalty2"}   # 迁移自 251229 penalty 家族


def norm_tag(tag: str) -> str:
    """归一化 tag：penalty*/punish → punish（D17 统一用 punish 做 9-tag 常量）。"""
    low = str(tag).lower()
    return "punish" if low.startswith("pena") or low == "punish" else str(tag)


# ---------------- 成绩源读取 ----------------

_NAME_HINTS = ("姓名", "name", "学生姓名", "学生")
_NUMBER_HINTS = ("学号", "number", "id", "学籍号", "student id", "studentid")


def _find_col(headers: list[str], hints: tuple[str, ...]) -> int | None:
    norm = [h.strip().lower() for h in headers]
    for i, h in enumerate(norm):
        if h in hints:
            return i
    for i, h in enumerate(norm):
        if h and any(token in h for token in hints if len(token) > 1):
            return i
    return None


def _s(row, i: int | None) -> str:
    if i is None or i >= len(row) or row[i] is None:
        return ""
    return str(row[i]).strip()


def read_score_xlsx(path: Path, col: str = "", weight: float = 1.0,
                    name_col_hint: str = "姓名") -> list[dict]:
    """从一张 xlsx 抽 {name, score, source, weight}：

    col 未指定/找不到时宽松回退：取"样本中数字占比最高"的非姓名列。"""
    path = Path(path)
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if len(rows) < 2:
        logger.warning(f"成绩表为空，跳过：{path}")
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    name_i = _find_col(headers, _NAME_HINTS)
    if name_i is None:
        name_i = next((i for i, h in enumerate(headers) if h.startswith(name_col_hint[:2])), 0)
    number_i = _find_col(headers, _NUMBER_HINTS)
    body = rows[1:]
    score_i = None
    if col:
        for i, h in enumerate(headers):
            if h and i != name_i and (col == h or col in h or col.lower() in h.lower()):
                score_i = i
                break
    if score_i is None:  # 宽松回退：数字占比最高的列
        best = (0, None)
        for i, h in enumerate(headers):
            if i == name_i or not h:
                continue
            nums = 0
            for r in body[:20]:
                v = r[i] if i < len(r) else None
                if v is None:
                    continue
                try:
                    float(v)
                    nums += 1
                except (TypeError, ValueError):
                    pass
            if nums > best[0]:
                best = (nums, i)
        score_i = best[1]
    if score_i is None:
        logger.warning(f"{path} 无法定位分数列（col={col}），跳过")
        return []
    out = []
    for r in body:
        if not r or r[name_i] is None:
            continue
        name = str(r[name_i]).strip()
        if not name:
            continue
        try:
            score = float(r[score_i])
        except (TypeError, ValueError):
            continue
        out.append({"name": name, "number": _s(r, number_i), "score": score,
                    "weight": float(weight),
                    "source": f"{path.name}:{col or headers[score_i]}"})
    logger.info(f"成绩源 {path.name} 读取 {len(out)} 条")
    return out


def _norm_id(value) -> str:
    """学号归一：去空格/常见分隔符、大小写不敏感（仅用于匹配回退）。"""
    return "".join(str(value or "").split()).lower()


def merge_scores(students: list[dict], score_rows: list[dict],
                 weight_normalize: bool = True) -> dict[str, dict]:
    """多源加权合并为每生 mean（迁移 _sortByScore 的均值排序语义；加权可配）。

    B2：优先按姓名匹配；姓名匹配不到时，若成绩行带 ``number`` 且点名册有学号，
    按学号回退匹配（返回结果仍 keyed by 点名册姓名，保持下游 tag_students 不变）。
    """
    by_name: dict[str, str] = {}
    by_number: dict[str, str] = {}
    for s in students:
        name = str(s.get("name", "")).strip()
        if not name:
            continue
        by_name[name] = name
        num = _norm_id(s.get("number", ""))
        if num:
            by_number.setdefault(num, name)
    buckets: dict[str, list[list[float]]] = {s["name"]: [] for s in students if s.get("name")}
    unmatched = 0
    for row in score_rows:
        raw_name = str(row.get("name", "")).strip()
        target = by_name.get(raw_name)
        if target is None:
            num = _norm_id(row.get("number", ""))
            target = by_number.get(num) if num else None
        if target is None:
            unmatched += 1
            continue
        buckets[target].append([row["weight"], row["score"]])
    if unmatched:
        logger.info(f"成绩匹配：{unmatched} 条未匹配到点名册（姓名/学号均未命中）")
    out: dict[str, dict] = {}
    for name, pairs in buckets.items():
        if not pairs:
            continue
        if weight_normalize:
            total_w = sum(p[0] for p in pairs)
            mean = sum(p[0] * p[1] for p in pairs) / total_w if total_w else 0.0
        else:
            mean = sum(p[1] for p in pairs) / len(pairs)
        out[name] = {"mean": round(mean, 2), "n": len(pairs)}
    return out


# ---------------- 分组打标签 ----------------

def tag_students(students: list[dict], scores: dict[str, dict],
                 group_cfg: list[dict] | None = None,
                 special_tag_cfg: dict[str, list[str]] | None = None) -> list[dict]:
    """按比例切分 tag（迁移 _sortByScore + _tag + _specialTag）。

    - 分数降序；无成绩学生排末尾；
    - 顺序切分各比例组（translation 与 punish 除外）；
    - 尾余续填最后一个 tag（迁移 rest 行为）；
    - translation 用等距选点重排（代替 np.random，可复现）；
    - special_tag_cfg: {tag: [名单]} 人工覆盖（最高优先级）。
    """
    group_cfg = group_cfg or DEFAULT_GROUP_CFG
    rows = [{"name": s["name"], "number": s.get("number", ""), "class": s.get("class", ""),
             "score": (scores.get(s["name"], {}) or {}).get("mean")} for s in students]
    rows.sort(key=lambda r: (r["score"] is None, -(r["score"] or 0)))
    n = len(rows)
    tag_list: list[str] = []
    rest = n
    for g in group_cfg:
        gname = norm_tag(g["group_name"])
        if gname in ("punish", RANDOM_GROUP):
            continue  # punish 不参与比例切分；translation 单独等距点选
        ratio = float(g["group_ratio"])
        take = int(n * ratio)
        tag_list.extend([gname] * take)
        rest -= take
    if rest > 0:
        last = tag_list[-1] if tag_list else "copy"
        tag_list.extend([last] * rest)
    elif rest < 0:
        raise ValueError(f"group_cfg 比例之和 > 1（rest={rest}）")
    tr = next((g for g in group_cfg if norm_tag(g["group_name"]) == RANDOM_GROUP), None)
    if tr and n > 0:
        k = min(n, int(n * float(tr["group_ratio"])))
        if k > 0:
            step = max(1, n // k)
            pos, cnt = k // 2 * step, 0
            while cnt < k and pos < n:
                tag_list[pos] = RANDOM_GROUP
                pos += step
                cnt += 1
    for r, tag in zip(rows, tag_list):
        r["tag"] = tag
    for tag, names in (special_tag_cfg or {}).items():
        for name in names:
            for r in rows:
                if r["name"] == name:
                    r.setdefault("old-tag", r.get("tag"))
                    r["tag"] = norm_tag(tag)
    return rows


def tagged_xlsx(path: Path, rows: list[dict]) -> Path:
    """写出带 tag 列的名单 xlsx（assist sheet make --roster 直接可用）。"""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "roster"
    ws.append(["姓名", "学号", "班级", "tag", "均值分", "old-tag"])
    for r in rows:
        ws.append([r["name"], r.get("number"), r.get("class"), r.get("tag"),
                   r.get("score"), ""] if r else [])
        if r.get("old-tag"):
            ws.cell(row=ws.max_row, column=6, value=", ".join(r["old-tag"]) if isinstance(r["old-tag"], list) else r["old-tag"])
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def tag_summary(rows: list[dict]) -> dict[str, int]:
    counts: dict[str, int] = {}
    for r in rows:
        counts[r["tag"]] = counts.get(r["tag"], 0) + 1
    return counts
